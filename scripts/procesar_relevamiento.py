#!/usr/bin/env python3
"""Genera los datos públicos del relevamiento sin modificar el Excel fuente."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


JURISDICTIONS = {
    "02": "Ciudad de Buenos Aires",
    "06": "Buenos Aires",
    "10": "Catamarca",
    "14": "Córdoba",
    "18": "Corrientes",
    "22": "Chaco",
    "26": "Chubut",
    "30": "Entre Ríos",
    "34": "Formosa",
    "38": "Jujuy",
    "42": "La Pampa",
    "46": "La Rioja",
    "50": "Mendoza",
    "54": "Misiones",
    "58": "Neuquén",
    "62": "Río Negro",
    "66": "Salta",
    "70": "San Juan",
    "74": "San Luis",
    "78": "Santa Cruz",
    "82": "Santa Fe",
    "86": "Santiago del Estero",
    "90": "Tucumán",
    "94": "Tierra del Fuego",
}


def normalize_cue(value: object) -> str:
    digits = re.sub(r"\D", "", str(value or ""))
    return digits.zfill(9) if digits else ""


def json_value(value: object) -> object:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def rows_as_dicts(sheet) -> tuple[list[str], list[dict[str, object]]]:
    values = list(sheet.iter_rows(values_only=True))
    if not values:
        return [], []
    headers = [str(value or "").strip() for value in values[0]]
    rows = []
    for row in values[1:]:
        if not any(value not in (None, "") for value in row):
            continue
        rows.append({header: json_value(row[index] if index < len(row) else None) for index, header in enumerate(headers)})
    return headers, rows


def find_header(headers: list[str], required_fragments: tuple[str, ...]) -> str:
    for header in headers:
        normalized = " ".join(header.split()).lower()
        if all(fragment.lower() in normalized for fragment in required_fragments):
            return header
    raise ValueError(f"No se encontró una columna con: {required_fragments}")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("excel", type=Path)
    parser.add_argument("geojson", type=Path)
    parser.add_argument("public_json", type=Path)
    parser.add_argument("validation_json", type=Path)
    parser.add_argument("--internal-json", type=Path)
    parser.add_argument("--weighted-base", type=Path)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    workbook = load_workbook(args.excel, data_only=True, read_only=True)
    required_sheets = {"CONFIGURACION", "MUESTRA", "CLASIFICACION", "RESPUESTAS", "PREGUNTAS"}
    missing_sheets = sorted(required_sheets - set(workbook.sheetnames))
    if missing_sheets:
        raise ValueError(f"Faltan hojas requeridas: {', '.join(missing_sheets)}")

    _, config_rows = rows_as_dicts(workbook["CONFIGURACION"])
    config = {str(row.get("Parámetro") or "").strip(): row.get("Valor") for row in config_rows if row.get("Parámetro")}
    universe = int(config.get("universo_operativo") or 1750)
    sample_total = int(config.get("muestra_proyectada") or 492)

    _, sample_rows = rows_as_dicts(workbook["MUESTRA"])
    sample_by_cue: dict[str, dict[str, str]] = {}
    jurisdiction_targets: dict[str, dict[str, object]] = {}
    duplicated_sample_cues: list[str] = []
    for row in sample_rows:
        jurisdiction_id = str(row.get("jurisdiccion_id") or "").zfill(2)
        if jurisdiction_id:
            jurisdiction_targets[jurisdiction_id] = {
                "id": jurisdiction_id,
                "name": str(row.get("jurisdiccion") or JURISDICTIONS.get(jurisdiction_id, jurisdiction_id)),
                "universe": int(row.get("universo_jurisdiccional") or 0),
                "target": int(row.get("muestra_jurisdiccional") or 0),
            }
        for field, role in (("cue_titular", "titular"), ("cue_reemplazo_1", "reemplazo"), ("cue_reemplazo_2", "reemplazo")):
            cue = normalize_cue(row.get(field))
            if not cue:
                continue
            if cue in sample_by_cue:
                duplicated_sample_cues.append(cue)
            sample_by_cue[cue] = {
                "sampleId": str(row.get("id_muestra") or ""),
                "role": role,
                "jurisdictionId": jurisdiction_id,
            }

    _, classification_rows = rows_as_dicts(workbook["CLASIFICACION"])
    classifications: dict[str, dict[str, str]] = {}
    for row in classification_rows:
        cue = normalize_cue(row.get("cue"))
        if cue:
            classifications[cue] = {
                "type": str(row.get("tipo_caso") or "").strip().lower(),
                "sampleId": str(row.get("id_muestra") or "").strip(),
                "note": str(row.get("observacion") or "").strip(),
            }

    weighted_rows: list[dict[str, object]] = []
    if args.weighted_base:
        with args.weighted_base.open(encoding="utf-8-sig", newline="") as source:
            reader = csv.DictReader(source)
            response_headers = list(reader.fieldnames or [])
            weighted_rows = [dict(row) for row in reader]
        required_weighted_headers = {"cue_norm", "ponderador_num", "estado_respuesta", "jurisdiction", "id", "rol"}
        missing_weighted_headers = sorted(required_weighted_headers - set(response_headers))
        if missing_weighted_headers:
            raise ValueError(f"Faltan columnas en la base ponderada: {', '.join(missing_weighted_headers)}")
        response_rows = [row for row in weighted_rows if str(row.get("estado_respuesta") or "").strip() == "Con respuesta"]
        cue_header = "cue_norm"
        sample_total = len(weighted_rows)
    else:
        response_headers, response_rows = rows_as_dicts(workbook["RESPUESTAS"])
        cue_header = find_header(response_headers, ("1.1.4. cue",))
    question_headers = {
        "q_2_1_1": find_header(response_headers, ("2.1.1.", "personal asignado formalmente")),
        "q_2_2_3": find_header(response_headers, ("2.2.3.", "remitido formalmente")),
        "q_2_2_5": find_header(response_headers, ("2.2.5.", "formato se estructuran")),
        "q_3_3": find_header(response_headers, ("3.3.", "sistematización")),
        "q_3_4": find_header(response_headers, ("3.4.", "registro actualizado")),
        "q_4_2": find_header(response_headers, ("4.2.", "vinculación con el sector")),
        "q_5_5": find_header(response_headers, ("5.5.", "oferentes o tutores externos")),
        "q_5_6": find_header(response_headers, ("5.6.", "capacidades profesionales")),
    }

    _, question_rows = rows_as_dicts(workbook["PREGUNTAS"])
    question_config = {
        str(row.get("id_pregunta") or ""): row
        for row in question_rows
        if str(row.get("id_pregunta") or "")
    }
    questions = []
    for order, (question_id, header) in enumerate(question_headers.items(), start=1):
        row = question_config.get(question_id, {})
        questions.append(
            {
                "id": question_id,
                "title": header or str(row.get("titulo_publico") or ""),
                "order": order,
            }
        )

    responses_by_cue: dict[str, list[dict[str, object]]] = defaultdict(list)
    empty_cue_rows = 0
    for row in response_rows:
        cue = normalize_cue(row.get(cue_header))
        if not cue:
            empty_cue_rows += 1
            continue
        responses_by_cue[cue].append(row)

    geo = json.loads(args.geojson.read_text(encoding="utf-8"))
    points_by_cue: dict[str, list[float]] = {}
    school_names_by_cue: dict[str, str] = {}
    for feature in geo.get("features", []):
        properties = feature.get("properties", {})
        cue = normalize_cue(properties.get("cue"))
        coordinates = feature.get("geometry", {}).get("coordinates")
        if cue and isinstance(coordinates, list) and len(coordinates) >= 2:
            points_by_cue[cue] = [float(coordinates[0]), float(coordinates[1])]
        school_name = str(properties.get("fna") or "").strip()
        if cue and school_name:
            school_names_by_cue[cue] = school_name

    duplicates = {cue: len(rows) for cue, rows in responses_by_cue.items() if len(rows) > 1}
    public_records = []
    internal_records = []
    private_issues = []
    covered_sample_ids: set[str] = set()
    covered_sample_ids_by_jurisdiction: dict[str, set[str]] = defaultdict(set)
    counts_by_jurisdiction: dict[str, Counter] = defaultdict(Counter)
    pending_classification = 0
    unmatched_geo = 0

    for index, cue in enumerate(sorted(responses_by_cue), start=1):
        row = responses_by_cue[cue][-1]
        jurisdiction_id = cue[:2]
        jurisdiction_name = JURISDICTIONS.get(jurisdiction_id, "Jurisdicción sin clasificar")
        sample_match = sample_by_cue.get(cue, {})
        manual = classifications.get(cue, {})
        if args.weighted_base:
            role = str(row.get("rol") or "").strip().lower()
            case_type = "reemplazo" if role.startswith("reemplazo") else "titular"
            sample_id = str(row.get("id") or "").strip()
            weight = float(str(row.get("ponderador_num") or "0").replace(",", "."))
        else:
            case_type = manual.get("type") or sample_match.get("role") or "muestra_provisional"
            sample_id = manual.get("sampleId") or sample_match.get("sampleId") or ""
            weight = 1.0
        if case_type == "muestra_provisional":
            pending_classification += 1
        if case_type != "complementaria" and sample_id:
            covered_sample_ids.add(sample_id)
            covered_sample_ids_by_jurisdiction[jurisdiction_id].add(sample_id)

        point = points_by_cue.get(cue)
        if not point:
            unmatched_geo += 1
            private_issues.append({"type": "cue_sin_geometria", "cue": cue})

        answers = {question_id: str(row.get(header) or "").strip() for question_id, header in question_headers.items()}
        public_id = f"escuela-{index:04d}"
        public_record = {
            "id": public_id,
            "jurisdictionId": jurisdiction_id,
            "jurisdiction": jurisdiction_name,
            "caseType": case_type,
            "weight": weight,
            "coordinates": point,
            "answers": answers,
        }
        public_records.append(public_record)
        internal_records.append({**public_record, "schoolName": school_names_by_cue.get(cue)})
        counts_by_jurisdiction[jurisdiction_id]["schools"] += 1
        if case_type == "complementaria":
            counts_by_jurisdiction[jurisdiction_id]["complementary"] += 1
        elif case_type == "reemplazo":
            counts_by_jurisdiction[jurisdiction_id]["replacements"] += 1
            counts_by_jurisdiction[jurisdiction_id]["sampleResponses"] += 1
        else:
            counts_by_jurisdiction[jurisdiction_id]["sampleResponses"] += 1

    jurisdictions = []
    mapping_complete = bool(args.weighted_base) or (len(sample_by_cue) >= sample_total and pending_classification == 0)
    provisional_covered_total = 0
    weighted_base_total = 0.0
    weighted_response_total = 0.0
    if args.weighted_base:
        jurisdiction_by_name = {name: code for code, name in JURISDICTIONS.items()}
        weighted_by_jurisdiction: dict[str, Counter] = defaultdict(Counter)
        for row in weighted_rows:
            jurisdiction_name = str(row.get("jurisdiction") or "").strip().replace("Santa Fé", "Santa Fe")
            jurisdiction_id = jurisdiction_by_name.get(jurisdiction_name)
            if not jurisdiction_id:
                raise ValueError(f"Jurisdicción sin código INDEC en la base ponderada: {jurisdiction_name}")
            weight = float(str(row.get("ponderador_num") or "0").replace(",", "."))
            stats = weighted_by_jurisdiction[jurisdiction_id]
            stats["sample"] += 1
            stats["weightTotal"] += weight
            if str(row.get("estado_respuesta") or "").strip() == "Con respuesta":
                stats["responses"] += 1
                stats["weightResponses"] += weight

        for jurisdiction_id in sorted(jurisdiction_targets):
            target = jurisdiction_targets[jurisdiction_id]
            counts = counts_by_jurisdiction[jurisdiction_id]
            stats = weighted_by_jurisdiction[jurisdiction_id]
            sample_count = int(stats["sample"])
            response_count = int(stats["responses"])
            weighted_base = float(stats["weightTotal"])
            weighted_responses = float(stats["weightResponses"])
            provisional_covered_total += response_count
            weighted_base_total += weighted_base
            weighted_response_total += weighted_responses
            jurisdictions.append(
                {
                    **target,
                    "target": sample_count,
                    "respondentSchools": response_count,
                    "sampleResponses": response_count,
                    "nonResponse": sample_count - response_count,
                    "covered": response_count,
                    "coverage": response_count / sample_count if sample_count else None,
                    "weightedBase": weighted_base,
                    "weightedResponses": weighted_responses,
                    "weightedNonResponse": weighted_base - weighted_responses,
                    "weightedCoverage": weighted_responses / weighted_base if weighted_base else None,
                    "replacements": int(counts["replacements"]),
                    "complementary": 0,
                }
            )
    else:
        for jurisdiction_id in sorted(jurisdiction_targets):
            target = jurisdiction_targets[jurisdiction_id]
            counts = counts_by_jurisdiction[jurisdiction_id]
            if mapping_complete:
                covered = len(covered_sample_ids_by_jurisdiction[jurisdiction_id])
            else:
                covered = min(int(counts["sampleResponses"]), int(target["target"]))
            provisional_covered_total += covered
            jurisdictions.append(
                {
                    **target,
                    "respondentSchools": int(counts["schools"]),
                    "sampleResponses": int(counts["sampleResponses"]),
                    "nonResponse": int(target["target"]) - covered,
                    "covered": covered,
                    "coverage": covered / int(target["target"]) if target["target"] else None,
                    "weightedBase": float(target["target"]),
                    "weightedResponses": float(covered),
                    "weightedNonResponse": float(int(target["target"]) - covered),
                    "weightedCoverage": covered / int(target["target"]) if target["target"] else None,
                    "replacements": int(counts["replacements"]),
                    "complementary": int(counts["complementary"]),
                }
            )
        weighted_base_total = float(sample_total)
        weighted_response_total = float(provisional_covered_total)

    source_path = args.weighted_base or args.excel
    source_hash = hashlib.sha256(source_path.read_bytes()).hexdigest()
    public_payload = {
        "metadata": {
            "generatedAt": datetime.now().astimezone().isoformat(timespec="seconds"),
            "sourceDate": config.get("fecha_actualizacion"),
            "sourceSha256": source_hash,
            "universe": universe,
            "sampleTarget": sample_total,
            "sampleShare": sample_total / universe if universe else None,
            "sourceRows": len(response_rows),
            "respondentSchools": len(public_records),
            "duplicateExtraRows": sum(value - 1 for value in duplicates.values()),
            "geolocatedSchools": len(public_records) - unmatched_geo,
            "unmatchedGeo": unmatched_geo,
            "coveredSamplePositions": len(covered_sample_ids) if mapping_complete else provisional_covered_total,
            "coverage": provisional_covered_total / sample_total if sample_total else None,
            "weightedBase": weighted_base_total,
            "weightedResponses": weighted_response_total,
            "weightedNonResponse": weighted_base_total - weighted_response_total,
            "weightedCoverage": weighted_response_total / weighted_base_total if weighted_base_total else None,
            "coverageStatus": "confirmada" if mapping_complete else "provisional",
            "pendingClassification": pending_classification,
            "sourceKind": "weighted_sample" if args.weighted_base else "responses_workbook",
            "access": "public",
        },
        "questions": questions,
        "jurisdictions": jurisdictions,
        "responses": public_records,
    }

    internal_payload = {
        **public_payload,
        "metadata": {**public_payload["metadata"], "access": "internal"},
        "responses": internal_records,
    }

    validation_payload = {
        "generatedAt": public_payload["metadata"]["generatedAt"],
        "sourceSha256": source_hash,
        "summary": {
            "sourceRows": len(response_rows),
            "uniqueCues": len(public_records),
            "emptyCueRows": empty_cue_rows,
            "duplicateCues": len(duplicates),
            "duplicateExtraRows": sum(value - 1 for value in duplicates.values()),
            "unmatchedGeo": unmatched_geo,
            "sampleCuesCompleted": len(sample_by_cue),
            "sampleTarget": sample_total,
            "pendingClassification": pending_classification,
            "coverageStatus": public_payload["metadata"]["coverageStatus"],
            "weightedBase": weighted_base_total,
            "weightedResponses": weighted_response_total,
            "weightedCoverage": public_payload["metadata"]["weightedCoverage"],
        },
        "duplicateDetails": [{"cue": cue, "rows": count} for cue, count in sorted(duplicates.items())],
        "issues": private_issues,
        "duplicatedSampleCues": sorted(set(duplicated_sample_cues)),
    }

    args.public_json.parent.mkdir(parents=True, exist_ok=True)
    args.validation_json.parent.mkdir(parents=True, exist_ok=True)
    args.public_json.write_text(json.dumps(public_payload, ensure_ascii=False, indent=2), encoding="utf-8")
    if args.internal_json:
        args.internal_json.parent.mkdir(parents=True, exist_ok=True)
        args.internal_json.write_text(json.dumps(internal_payload, ensure_ascii=False, indent=2), encoding="utf-8")
    args.validation_json.write_text(json.dumps(validation_payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(validation_payload["summary"], ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
